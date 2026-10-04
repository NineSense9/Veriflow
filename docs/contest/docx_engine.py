"""Contest docx: Songti body, Times New Roman Latin, Word TOC field."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor, Twips
from docx.enum.style import WD_STYLE_TYPE

A4_W = Cm(21.0)
A4_H = Cm(29.7)
LEFT = Cm(2.8)
RIGHT = Cm(2.5)
TOP = Cm(2.5)
BOTTOM = Cm(2.4)
CONTENT_WIDTH = Cm(15.7)

NAVY = "1F4E79"
HEADER_LINE = "1F4E79"
ZEBRA = "F3F6FA"
HEAD_BG = "E7EEF6"


def _rfonts(rPr, east="宋体", latin="Times New Roman"):
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for key in list(rFonts.attrib):
        if key.rsplit("}", 1)[-1].lower().endswith("theme"):
            del rFonts.attrib[key]
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:cs"), latin)
    rFonts.set(qn("w:eastAsia"), east)


def set_run_font(run, size_pt, bold=False, latin="Times New Roman", east="宋体"):
    run.bold = bold
    run.font.size = Pt(size_pt)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run.font.name = latin
    _rfonts(run._element.get_or_add_rPr(), east=east, latin=latin)


def set_style_font(style, size_pt, bold=False, latin="Times New Roman", east="宋体"):
    style.font.name = latin
    style.font.size = Pt(size_pt)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    _rfonts(style.element.get_or_add_rPr(), east=east, latin=latin)


def set_indent_chars(pPr, chars):
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    if chars:
        ind.set(qn("w:firstLineChars"), str(chars))
        ind.set(qn("w:firstLine"), str(int(chars / 100 * 240)))
    else:
        ind.set(qn("w:firstLineChars"), "0")
        ind.set(qn("w:firstLine"), "0")


def style_indent(style, chars):
    set_indent_chars(style.element.get_or_add_pPr(), chars)


def paragraph_indent(paragraph, chars):
    set_indent_chars(paragraph._p.get_or_add_pPr(), chars)


def shade_cell(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.find(qn("w:tcBorders"))
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        el = tcBorders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            tcBorders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "B0B8C4")


def set_table_width(table, widths):
    table.autofit = False
    table.allow_autofit = False
    tbl = table._tbl
    tblPr = tbl.tblPr
    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")
    cell_mar = tblPr.find(qn("w:tblCellMar"))
    if cell_mar is None:
        cell_mar = OxmlElement("w:tblCellMar")
        tblPr.append(cell_mar)
    for edge, size in (("top", "40"), ("left", "80"), ("bottom", "40"), ("right", "80")):
        el = cell_mar.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            cell_mar.append(el)
        el.set(qn("w:w"), size)
        el.set(qn("w:type"), "dxa")
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = sum(widths)
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")
    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        for child in list(grid):
            grid.remove(child)
    else:
        grid = OxmlElement("w:tblGrid")
        tblPr.addnext(grid)
    for width in widths:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(width))
        grid.append(gc)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(width))
            tcW.set(qn("w:type"), "dxa")


def prevent_row_split(row):
    tr = row._tr
    trPr = tr.get_or_add_trPr()
    cant = OxmlElement("w:cantSplit")
    trPr.append(cant)


def add_page_field(paragraph, kind="PAGE"):
    run = paragraph.add_run()
    set_run_font(run, 9)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {kind} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r = run._r
    r.append(begin)
    r.append(instr)
    r.append(separate)
    shown = OxmlElement("w:t")
    shown.text = "1"
    r.append(shown)
    r.append(end)


def add_bottom_border(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), HEADER_LINE)
    pBdr.append(bottom)
    pPr.append(pBdr)


def configure_styles(doc):
    for style in doc.styles:
        _rfonts(style.element.get_or_add_rPr())
    defaults = doc.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        for rPr in defaults.iter(qn("w:rPr")):
            _rfonts(rPr)
    normal = doc.styles["Normal"]
    set_style_font(normal, 12, False)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    style_indent(normal, 200)

    specs = [
        ("Heading 1", 16, True, 0, 16, 8),
        ("Heading 2", 14, True, 1, 12, 6),
        ("Heading 3", 12, True, 2, 8, 4),
    ]
    for name, size, bold, level, before, after in specs:
        style = doc.styles[name]
        set_style_font(style, size, bold)
        pf = style.paragraph_format
        pf.line_spacing_rule = WD_LINE_SPACING.SINGLE
        pf.space_before = Pt(before)
        pf.space_after = Pt(after)
        pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        pf.keep_with_next = True
        style_indent(style, 0)
        pPr = style.element.get_or_add_pPr()
        outline = pPr.find(qn("w:outlineLvl"))
        if outline is None:
            outline = OxmlElement("w:outlineLvl")
            pPr.append(outline)
        outline.set(qn("w:val"), str(level))

    for level, size, left in ((1, 12, 0), (2, 12, 420), (3, 12, 840)):
        name = f"TOC {level}"
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        set_style_font(style, size, level == 1)
        pf = style.paragraph_format
        pf.line_spacing = 1.5
        pf.space_before = Pt(2)
        pf.space_after = Pt(2)
        pf.left_indent = Twips(left)
        style_indent(style, 0)
        pf.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)

    for name, size, bold, center in (
        ("CoverKicker", 14, False, True),
        ("CoverName", 28, True, True),
        ("CoverTitle", 22, True, True),
        ("CoverSub", 12, False, True),
        ("TocTitle", 18, True, True),
        ("Caption", 10.5, True, True),
        ("TableText", 10.5, False, False),
        ("CodeText", 10.5, False, False),
        ("MetaLine", 12, False, True),
    ):
        try:
            style = doc.styles[name]
        except KeyError:
            style = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        latin = "Times New Roman"
        set_style_font(style, size, bold, latin=latin)
        pf = style.paragraph_format
        pf.line_spacing = 1.5 if name not in {"TableText", "CodeText", "Caption"} else 1.0
        pf.space_before = Pt(0)
        pf.space_after = Pt(0)
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        style_indent(style, 0)
        if name == "CodeText":
            pf.left_indent = Cm(0.75)


def mark_update_fields(doc):
    settings = doc.settings.element
    node = settings.find(qn("w:updateFields"))
    if node is None:
        node = OxmlElement("w:updateFields")
        settings.append(node)
    node.set(qn("w:val"), "true")


def add_toc(paragraph):
    """Real Word TOC field. Page numbers fill when Word updates fields."""
    run = paragraph.add_run()
    r = run._r
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = r' TOC \o "1-3" \h \z \u '
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    r.append(begin)
    r.append(instr)
    r.append(separate)
    placeholder = paragraph.add_run("目录将在打开文档后按标题自动生成页码。")
    set_run_font(placeholder, 12)
    end_run = paragraph.add_run()
    end_run._r.append(end)


def header_footer(section, header_text):
    section.different_first_page_header_footer = True
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    section.first_page_header.is_linked_to_previous = False
    section.first_page_footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph_indent(hp, 0)
    hp.paragraph_format.space_after = Pt(2)
    run = hp.add_run(header_text)
    set_run_font(run, 9)
    add_bottom_border(hp)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_indent(fp, 0)
    left = fp.add_run("第 ")
    set_run_font(left, 9)
    add_page_field(fp, "PAGE")
    right = fp.add_run(" 页")
    set_run_font(right, 9)
    # Keep the cover free of header and page number.
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    run = p.add_run(text)
    set_run_font(run, 12)
    paragraph_indent(p, 200)
    return p


def add_heading(doc, text, level):
    style = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3"}[level]
    size = {1: 16, 2: 14, 3: 12}[level]
    p = doc.add_paragraph(style=style)
    run = p.add_run(text)
    set_run_font(run, size, bold=True)
    paragraph_indent(p, 0)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    run = p.add_run(text)
    set_run_font(run, 10.5, bold=True)
    paragraph_indent(p, 0)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)
    return p


def add_code(doc, text):
    for line in text.split("\n"):
        p = doc.add_paragraph(style="CodeText")
        run = p.add_run(line if line else " ")
        set_run_font(run, 10.5)
        paragraph_indent(p, 0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)


def _visual_len(text):
    total = 0.0
    for ch in str(text):
        total += 1.0 if ord(ch) > 127 else 0.55
    return total


def _twips_for(visual):
    # 10.5pt Songti: one CJK em is 210 twips. ASCII is already scaled in visual.
    # 160 twips of cell margin, plus slack so the last glyph is not pushed to the next line.
    return int(visual * 210 + 480)


def _column_widths(headers, rows):
    """Fit short labels on one line. Long sentences may wrap. Widths sum to the text area."""
    usable = 8902
    count = len(headers)
    visuals = []
    for index, header in enumerate(headers):
        column = [_visual_len(header)]
        for row in rows:
            column.append(_visual_len(row[index]))
        visuals.append(column)

    def mins_for(limit):
        mins = []
        for column in visuals:
            chosen = column[0]
            for visual in column:
                if visual <= limit:
                    chosen = max(chosen, visual)
            mins.append(max(900, _twips_for(chosen)))
        return mins

    minimums = mins_for(0)
    for limit in (8, 10, 12, 14, 16, 18, 22):
        candidate = mins_for(limit)
        if sum(candidate) <= usable:
            minimums = candidate
        else:
            break
    bulk = [max(column) for column in visuals]
    total = sum(bulk) or 1
    widths = [max(minimums[index], int(usable * bulk[index] / total)) for index in range(count)]
    overflow = sum(widths) - usable
    guard = 0
    while overflow > 0 and guard < 16:
        donor = max(range(count), key=lambda index: widths[index] - minimums[index])
        room = widths[donor] - minimums[donor]
        if room <= 0:
            break
        cut = min(overflow, room)
        widths[donor] -= cut
        overflow -= cut
        guard += 1
    if min(widths) <= 0:
        base = usable // count
        widths = [base] * count
    drift = usable - sum(widths)
    widths[max(range(count), key=lambda index: widths[index])] += drift
    return widths


def add_table(doc, headers, rows):
    count = len(headers)
    table = doc.add_table(rows=1 + len(rows), cols=count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    widths = _column_widths(headers, rows)
    set_table_width(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        p = cell.paragraphs[0]
        p.style = doc.styles["TableText"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph_indent(p, 0)
        run = p.add_run(header)
        set_run_font(run, 10.5, bold=True)
        shade_cell(cell, HEAD_BG)
        set_cell_border(cell)
    prevent_row_split(table.rows[0])
    header_marker = OxmlElement("w:tblHeader")
    header_marker.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(header_marker)
    for r_index, row in enumerate(rows):
        for c_index, value in enumerate(row):
            cell = table.rows[r_index + 1].cells[c_index]
            cell.text = ""
            p = cell.paragraphs[0]
            p.style = doc.styles["TableText"]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph_indent(p, 0)
            run = p.add_run(str(value))
            set_run_font(run, 10.5, bold=False)
            if r_index % 2 == 1:
                shade_cell(cell, ZEBRA)
            set_cell_border(cell)
        prevent_row_split(table.rows[r_index + 1])
    spacer = doc.add_paragraph(style="Caption")
    spacer.paragraph_format.space_after = Pt(6)
    return table


def build_document(meta, blocks, path: Path):
    doc = Document()
    configure_styles(doc)
    mark_update_fields(doc)
    section = doc.sections[0]
    section.page_width = A4_W
    section.page_height = A4_H
    section.left_margin = LEFT
    section.right_margin = RIGHT
    section.top_margin = TOP
    section.bottom_margin = BOTTOM
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.15)
    header_footer(section, meta["header"])

    def centered(style, text, size, bold=False, before=0, after=0):
        p = doc.add_paragraph(style=style)
        paragraph_indent(p, 0)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        run = p.add_run(text)
        set_run_font(run, size, bold=bold)
        return p

    centered("CoverKicker", meta["kicker"], 14, before=72, after=18)
    centered("CoverName", "VeriFlow", 28, bold=True, before=18, after=12)
    centered("CoverTitle", meta["doc_title"], 22, bold=True, before=6, after=12)
    centered("CoverSub", meta["subtitle"], 12, before=6, after=18)
    for line in meta["meta_lines"]:
        centered("MetaLine", line, 12, before=2, after=2)
    centered("MetaLine", "编制日期：2026年09月  核对更新：2026年10月05日", 12, before=18, after=0)

    doc.add_page_break()
    toc_title = doc.add_paragraph(style="TocTitle")
    paragraph_indent(toc_title, 0)
    toc_run = toc_title.add_run("目录")
    set_run_font(toc_run, 18, bold=True)
    toc_p = doc.add_paragraph()
    paragraph_indent(toc_p, 0)
    add_toc(toc_p)
    doc.add_page_break()

    for block in blocks:
        kind = block[0]
        if kind == "h1":
            add_heading(doc, block[1], 1)
        elif kind == "h2":
            add_heading(doc, block[1], 2)
        elif kind == "h3":
            add_heading(doc, block[1], 3)
        elif kind == "p":
            add_body(doc, block[1])
        elif kind == "cap":
            add_caption(doc, block[1])
        elif kind == "table":
            add_table(doc, block[1], block[2])
        elif kind == "code":
            add_code(doc, block[1])
        else:
            raise ValueError(kind)

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    return path
